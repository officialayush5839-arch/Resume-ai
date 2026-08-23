import docx
doc = docx.Document()
doc.add_paragraph('John Doe')
doc.add_paragraph('john.doe.123456@vupune.ac.in')
doc.add_paragraph('Experience: 5 years at Google. Skills: Python, React, SQL.')
doc.save('test_resume.docx')
